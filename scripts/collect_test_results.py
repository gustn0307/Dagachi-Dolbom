from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

try:
    from openpyxl import load_workbook
except ImportError:
    print(
        "[ERROR] openpyxl이 설치되어 있지 않습니다.\n"
        "CMD에서 다음 명령을 실행하세요:\n"
        "python -m pip install openpyxl",
        file=sys.stderr,
    )
    sys.exit(1)


REQ_PATTERN = re.compile(r"REQ-[A-Z]+-\d{2,3}", re.IGNORECASE)
PYTEST_REQ_PATTERN = re.compile(r"(?:REQ|req)[_-]([A-Za-z]+)[_-](\d{2,3})")


@dataclass
class TestResult:
    source: str
    layer: str
    classname: str
    name: str
    status: str
    duration: float
    req_ids: list[str]
    evidence: str


def normalize_req_id(raw: str) -> str:
    return raw.upper().replace("_", "-")


def extract_req_ids(*texts: str) -> list[str]:
    combined = " ".join(text for text in texts if text)
    found: set[str] = set()

    # REQ-AI-12 형태
    for match in REQ_PATTERN.findall(combined):
        found.add(normalize_req_id(match))

    # pytest의 req_ai_12 / REQ_AI_12 형태
    for domain, number in PYTEST_REQ_PATTERN.findall(combined):
        found.add(f"REQ-{domain.upper()}-{number}")

    return sorted(found)


def testcase_status(testcase: ET.Element) -> str:
    if testcase.find("failure") is not None or testcase.find("error") is not None:
        return "Fail"
    if testcase.find("skipped") is not None:
        return "Skip"
    return "Pass"


def parse_junit_file(
    path: Path,
    source: str,
    layer: str,
) -> list[TestResult]:
    root = ET.parse(path).getroot()
    results: list[TestResult] = []

    for testcase in root.iter("testcase"):
        classname = testcase.attrib.get("classname", "").strip()
        name = testcase.attrib.get("name", "").strip()

        try:
            duration = float(testcase.attrib.get("time", "0") or 0)
        except ValueError:
            duration = 0.0

        results.append(
            TestResult(
                source=source,
                layer=layer,
                classname=classname,
                name=name,
                status=testcase_status(testcase),
                duration=duration,
                req_ids=extract_req_ids(classname, name),
                evidence=str(path),
            )
        )

    return results


def collect_backend(root: Path) -> list[TestResult]:
    result_dir = root / "backend" / "build" / "test-results" / "test"
    files = sorted(result_dir.glob("*.xml"))

    if not files:
        print(f"[WARN] Backend XML 없음: {result_dir}")
        return []

    results: list[TestResult] = []
    for path in files:
        results.extend(parse_junit_file(path, "Backend", "BE"))

    return results


def collect_ai(root: Path) -> list[TestResult]:
    path = root / "ai-service" / "test-results" / "pytest-junit.xml"

    if not path.exists():
        print(f"[WARN] AI XML 없음: {path}")
        return []

    return parse_junit_file(path, "AI Service", "AI")


def collect_frontend(root: Path) -> list[TestResult]:
    path = root / "frontend" / "junit-results" / "playwright-junit.xml"

    if not path.exists():
        print(f"[WARN] Frontend XML 없음: {path}")
        return []

    return parse_junit_file(path, "Frontend", "FE")


def collect_all(root: Path) -> list[TestResult]:
    return (
        collect_backend(root)
        + collect_ai(root)
        + collect_frontend(root)
    )


def load_requirement_map(ws) -> dict[str, dict]:
    result: dict[str, dict] = {}

    for row in range(2, ws.max_row + 1):
        req_id = ws.cell(row=row, column=2).value

        if not req_id:
            continue

        req_id = str(req_id).strip().upper()

        result[req_id] = {
            "no": ws.cell(row=row, column=1).value,
            "feature": ws.cell(row=row, column=4).value,
            "priority": ws.cell(row=row, column=7).value,
            "owner": ws.cell(row=row, column=9).value,
        }

    return result


def clear_test_results(ws) -> None:
    # Header(row 1)는 보존
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)


def append_test_result_rows(
    ws,
    tests: Iterable[TestResult],
    requirements: dict[str, dict],
) -> tuple[int, int]:
    mapped_count = 0
    unmapped_count = 0

    for test in tests:
        if not test.req_ids:
            unmapped_count += 1

            ws.append(
                [
                    f"AUTO-UNMAPPED-{unmapped_count:04d}",
                    "UNMAPPED",
                    None,
                    None,
                    None,
                    test.layer,
                    "자동/JUnit XML",
                    f"{test.classname} :: {test.name}",
                    test.status,
                    test.classname,
                    test.evidence,
                    None,
                    date.today().isoformat(),
                    None,
                    "해당없음",
                    None,
                    f"{test.source} / REQ-ID 미매핑 / {test.duration:.3f}s",
                ]
            )
            continue

        for req_id in test.req_ids:
            req = requirements.get(req_id)

            if req is None:
                unmapped_count += 1

                ws.append(
                    [
                        f"AUTO-UNKNOWN-{unmapped_count:04d}",
                        req_id,
                        None,
                        None,
                        None,
                        test.layer,
                        "자동/JUnit XML",
                        f"{test.classname} :: {test.name}",
                        test.status,
                        test.classname,
                        test.evidence,
                        None,
                        date.today().isoformat(),
                        None,
                        "해당없음",
                        None,
                        f"{test.source} / Requirements에 없는 REQ-ID",
                    ]
                )
                continue

            mapped_count += 1

            ws.append(
                [
                    f"AUTO-{req_id}-{mapped_count:04d}",
                    req_id,
                    req["no"],
                    req["feature"],
                    req["priority"],
                    test.layer,
                    "자동/JUnit XML",
                    f"{test.classname} :: {test.name}",
                    test.status,
                    test.classname,
                    test.evidence,
                    req["owner"],
                    date.today().isoformat(),
                    None,
                    "해당없음",
                    None,
                    f"{test.source} / {test.duration:.3f}s",
                ]
            )

    return mapped_count, unmapped_count


def append_excluded_placeholders(
    ws,
    requirements: dict[str, dict],
) -> int:
    count = 0

    for req_id, req in requirements.items():
        if req["priority"] != "제외":
            continue

        count += 1

        ws.append(
            [
                f"OUT-{req_id}",
                req_id,
                req["no"],
                req["feature"],
                req["priority"],
                None,
                "Release Scope",
                "Final Release Scope 제외",
                "제외",
                None,
                None,
                req["owner"],
                None,
                None,
                "해당없음",
                None,
                "OUT-RELEASE",
            ]
        )

    return count


def export_raw_results(
    output_dir: Path,
    tests: list[TestResult],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)

    json_path = output_dir / "collected-test-results.json"
    csv_path = output_dir / "collected-test-results.csv"

    with json_path.open("w", encoding="utf-8") as fp:
        json.dump(
            [
                {
                    "source": t.source,
                    "layer": t.layer,
                    "classname": t.classname,
                    "name": t.name,
                    "status": t.status,
                    "duration": t.duration,
                    "req_ids": t.req_ids,
                    "evidence": t.evidence,
                }
                for t in tests
            ],
            fp,
            ensure_ascii=False,
            indent=2,
        )

    with csv_path.open("w", encoding="utf-8-sig", newline="") as fp:
        writer = csv.writer(fp)
        writer.writerow(
            [
                "source",
                "layer",
                "classname",
                "test_name",
                "status",
                "duration",
                "req_ids",
                "evidence",
            ]
        )

        for t in tests:
            writer.writerow(
                [
                    t.source,
                    t.layer,
                    t.classname,
                    t.name,
                    t.status,
                    f"{t.duration:.3f}",
                    ",".join(t.req_ids),
                    t.evidence,
                ]
            )


def print_summary(tests: list[TestResult]) -> None:
    total = len(tests)
    passed = sum(t.status == "Pass" for t in tests)
    failed = sum(t.status == "Fail" for t in tests)
    skipped = sum(t.status == "Skip" for t in tests)
    mapped_tests = sum(bool(t.req_ids) for t in tests)
    unmapped_tests = total - mapped_tests

    print()
    print("========================================")
    print("Test Result Collector")
    print("========================================")
    print(f"전체 testcase : {total}")
    print(f"Pass          : {passed}")
    print(f"Fail          : {failed}")
    print(f"Skip          : {skipped}")
    print(f"REQ 매핑 테스트: {mapped_tests}")
    print(f"UNMAPPED      : {unmapped_tests}")
    print("========================================")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Dagachi-Dolbom JUnit/pytest/Playwright 결과를 Excel에 수집합니다."
    )

    parser.add_argument(
        "--project-root",
        default=None,
        help="Dagachi-Dolbom 프로젝트 루트. 기본값은 스크립트 기준 상위 폴더",
    )

    parser.add_argument(
        "--excel",
        required=True,
        help="v1.4.1 Excel 템플릿 경로",
    )

    parser.add_argument(
        "--output",
        default=None,
        help="생성할 Excel 경로. 생략하면 입력 파일명 뒤에 _collected 추가",
    )

    args = parser.parse_args()

    script_path = Path(__file__).resolve()
    project_root = (
        Path(args.project_root).resolve()
        if args.project_root
        else script_path.parent.parent
    )

    excel_path = Path(args.excel).resolve()

    if not excel_path.exists():
        raise FileNotFoundError(f"Excel 파일을 찾을 수 없습니다: {excel_path}")

    if args.output:
        output_path = Path(args.output).resolve()
    else:
        output_path = excel_path.with_name(
            f"{excel_path.stem}_collected{excel_path.suffix}"
        )

    tests = collect_all(project_root)

    print_summary(tests)

    raw_output_dir = project_root / "test-report-output"
    export_raw_results(raw_output_dir, tests)

    wb = load_workbook(excel_path)

    if "Requirements" not in wb.sheetnames:
        raise RuntimeError("Excel에 Requirements 시트가 없습니다.")

    if "Test Results" not in wb.sheetnames:
        raise RuntimeError("Excel에 Test Results 시트가 없습니다.")

    requirements_ws = wb["Requirements"]
    results_ws = wb["Test Results"]

    requirements = load_requirement_map(requirements_ws)

    clear_test_results(results_ws)

    mapped_rows, unmapped_rows = append_test_result_rows(
        results_ws,
        tests,
        requirements,
    )

    excluded_rows = append_excluded_placeholders(
        results_ws,
        requirements,
    )

    wb.save(output_path)

    print()
    print(f"[OK] Excel 생성: {output_path}")
    print(f"[OK] RAW JSON/CSV: {raw_output_dir}")
    print(f"[OK] REQ 매핑 행: {mapped_rows}")
    print(f"[OK] UNMAPPED/UNKNOWN 행: {unmapped_rows}")
    print(f"[OK] OUT-RELEASE 행: {excluded_rows}")


if __name__ == "__main__":
    main()